from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font


FORMATO_MONEDA = '#,##0.00'
FORMATO_FECHA = 'yyyy-mm-dd'


def _texto_seguro_excel(valor):
    if valor is None:
        return ""

    texto = str(valor)

    if texto.startswith(("=", "+", "-", "@")):
        return f"'{texto}"

    return texto


def _ajustar_columnas(hoja):
    for columna in hoja.columns:
        largo_maximo = 0
        letra = columna[0].column_letter

        for celda in columna:
            if celda.value is None:
                continue

            largo_maximo = max(
                largo_maximo,
                len(str(celda.value)),
            )

        hoja.column_dimensions[letra].width = min(
            max(largo_maximo + 2, 12),
            40,
        )


def _estilizar_encabezado(hoja):
    for celda in hoja[1]:
        celda.font = Font(bold=True)

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions


def generar_excel_pagos(*, pagos):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Pagos"

    encabezados = [
        "Fecha",
        "Alumno",
        "Identificación",
        "Medio de pago",
        "Referencia",
        "Estado",
        "Moneda",
        "Valor",
        "Aplicado",
        "Saldo disponible",
    ]

    hoja.append(encabezados)

    for fila in pagos:
        hoja.append(
            [
                fila.fecha_pago,
                _texto_seguro_excel(fila.nombre_alumno),
                _texto_seguro_excel(
                    fila.numero_identidad
                ),
                fila.medio_pago,
                _texto_seguro_excel(
                    fila.referencia
                ),
                fila.estado,
                fila.moneda,
                float(fila.valor),
                float(fila.total_aplicado),
                float(fila.saldo_sin_aplicar),
            ]
        )

    _estilizar_encabezado(hoja)

    for celda in hoja["A"][1:]:
        celda.number_format = FORMATO_FECHA

    for columna in ("H", "I", "J"):
        for celda in hoja[columna][1:]:
            celda.number_format = FORMATO_MONEDA

    _ajustar_columnas(hoja)

    salida = BytesIO()
    libro.save(salida)
    salida.seek(0)

    return salida


def generar_excel_cartera(
    *,
    resumen,
    alumnos,
):
    libro = Workbook()

    hoja_resumen = libro.active
    hoja_resumen.title = "Resumen"

    hoja_resumen.append(
        [
            "Indicador",
            "Valor",
        ]
    )

    indicadores = [
        (
            "Total obligaciones",
            resumen.total_obligaciones,
            True,
        ),
        (
            "Total aplicado",
            resumen.total_aplicado,
            True,
        ),
        (
            "Saldo pendiente",
            resumen.saldo_pendiente,
            True,
        ),
        (
            "Saldo vigente",
            resumen.saldo_vigente,
            True,
        ),
        (
            "Saldo vencido",
            resumen.saldo_vencido,
            True,
        ),
        (
            "Saldo disponible",
            resumen.saldo_pagos_sin_aplicar,
            True,
        ),
        (
            "Alumnos con saldo",
            resumen.cantidad_alumnos_con_saldo,
            False,
        ),
        (
            "Alumnos con saldo vencido",
            resumen.cantidad_alumnos_con_saldo_vencido,
            False,
        ),
        (
            "Obligaciones pendientes",
            resumen.cantidad_obligaciones_pendientes,
            False,
        ),
        (
            "Obligaciones parciales",
            resumen.cantidad_obligaciones_parciales,
            False,
        ),
        (
            "Obligaciones pagadas",
            resumen.cantidad_obligaciones_pagadas,
            False,
        ),
        (
            "Obligaciones vencidas",
            resumen.cantidad_obligaciones_vencidas,
            False,
        ),
        (
            "Vence hoy",
            resumen.saldo_vence_hoy,
            True,
        ),
        (
            "1 - 30 días",
            resumen.saldo_1_30,
            True,
        ),
        (
            "31 - 60 días",
            resumen.saldo_31_60,
            True,
        ),
        (
            "61 - 90 días",
            resumen.saldo_61_90,
            True,
        ),
        (
            "Más de 90 días",
            resumen.saldo_mas_90,
            True,
        ),
        (
            "Sin vencimiento",
            resumen.saldo_sin_vencimiento,
            True,
        ),
    ]

    for nombre, valor, es_moneda in indicadores:
        hoja_resumen.append(
            [
                nombre,
                float(valor) if es_moneda else valor,
            ]
        )

        if es_moneda:
            hoja_resumen.cell(
                row=hoja_resumen.max_row,
                column=2,
            ).number_format = FORMATO_MONEDA

    for celda in hoja_resumen[1]:
        celda.font = Font(bold=True)

    _ajustar_columnas(hoja_resumen)

    hoja_cartera = libro.create_sheet(
        "Cartera"
    )

    encabezados = [
        "Alumno",
        "Total obligaciones",
        "Total aplicado",
        "Saldo pendiente",
        "Saldo vigente",
        "Saldo vencido",
        "Días atraso",
        "Saldo disponible",
        "Pendientes",
        "Parciales",
        "Vencidas",
        "Último pago",
    ]

    hoja_cartera.append(encabezados)

    for fila in alumnos:
        dias_atraso = (
            fila.dias_atraso_max
            if fila.cantidad_obligaciones_vencidas
            else None
        )

        hoja_cartera.append(
            [
                _texto_seguro_excel(
                    fila.nombre
                ),
                float(
                    fila.total_obligaciones
                ),
                float(
                    fila.total_aplicado
                ),
                float(
                    fila.saldo_pendiente
                ),
                float(
                    fila.saldo_vigente
                ),
                float(
                    fila.saldo_vencido
                ),
                dias_atraso,
                float(
                    fila.saldo_pagos_sin_aplicar
                ),
                fila.cantidad_obligaciones_pendientes,
                fila.cantidad_obligaciones_parciales,
                fila.cantidad_obligaciones_vencidas,
                fila.ultima_fecha_pago,
            ]
        )

    _estilizar_encabezado(hoja_cartera)

    for columna in (
        "B",
        "C",
        "D",
        "E",
        "F",
        "H",
    ):
        for celda in hoja_cartera[columna][1:]:
            celda.number_format = FORMATO_MONEDA

    for celda in hoja_cartera["L"][1:]:
        celda.number_format = FORMATO_FECHA

    _ajustar_columnas(hoja_cartera)

    salida = BytesIO()
    libro.save(salida)
    salida.seek(0)

    return salida