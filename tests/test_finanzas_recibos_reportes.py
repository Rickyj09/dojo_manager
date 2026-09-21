from datetime import date
from decimal import Decimal

from app.extensions import db as _db
from app.models.finanzas import ObligacionFinanciera
from app.services.finanzas.pagos import (
    anular_pago,
    aplicar_pago_a_obligaciones,
    registrar_pago,
)
from app.services.finanzas.reportes import (
    obtener_reporte_pagos,
)
from app.models.sucursal import Sucursal
from io import BytesIO

from openpyxl import load_workbook
from app.models.role import Role
def login(client, user):
    response = client.post(
        "/auth/login",
        data={
            "username": user.username,
            "password": "secret",
        },
    )
    assert response.status_code == 302

def crear_pago(
    base_data,
    *,
    academia_key="academia_a",
    alumno_key="alumno_a1",
    valor="50.00",
    medio_pago="EFECTIVO",
    referencia=None,
    observacion=None,
    fecha_pago=date(2026, 9, 5),
):

    pago = registrar_pago(
        academia_id=base_data[academia_key].id,
        alumno_id=base_data[alumno_key].id,
        fecha_pago=fecha_pago,
        valor=Decimal(valor),
        medio_pago=medio_pago,
        moneda="USD",
        referencia=referencia,
        observacion=observacion,
    )

    _db.session.commit()
    return pago

def crear_obligacion(
    base_data,
    *,
    academia_key="academia_a",
    alumno_key="alumno_a1",
    valor="50.00",
    periodo="2026-09",
    concepto="Pension septiembre",
):
    valor_decimal = Decimal(valor)

    obligacion = ObligacionFinanciera(
        academia_id=base_data[academia_key].id,
        alumno_id=base_data[alumno_key].id,
        alumno_plan_financiero_id=None,
        periodo=periodo,
        tipo_obligacion="PENSION",
        concepto=concepto,
        origen="TEST_RECIBO",
        fecha_emision=date(2026, 9, 1),
        fecha_vencimiento=date(2026, 9, 7),
        tarifa_base_snapshot=valor_decimal,
        porcentaje_descuento_snapshot=None,
        valor_descuento_snapshot=Decimal("0.00"),
        valor_final_snapshot=valor_decimal,
        moneda_snapshot="USD",
        estado="PENDIENTE",
    )

    _db.session.add(obligacion)
    _db.session.commit()

    return obligacion


def test_admin_puede_ver_recibo(app, db, base_data):
    pago = crear_pago(
        base_data,
        valor="75.50",
        medio_pago="TRANSFERENCIA",
        referencia="TRX-001",
        observacion="Pago mensual",
    )

    client = app.test_client()
    login(client, base_data["admin_a"])

    response = client.get(
        f"/finanzas/pagos/{pago.id}/recibo"
    )

    assert response.status_code == 200
    assert b"Recibo interno de pago" in response.data

    numero = f"REC-{pago.id:06d}".encode()
    assert numero in response.data

    assert base_data["academia_a"].nombre.encode() in response.data
    assert base_data["alumno_a1"].nombres.encode() in response.data
    assert base_data["alumno_a1"].apellidos.encode() in response.data

    assert b"TRANSFERENCIA" in response.data
    assert b"TRX-001" in response.data
    assert b"Pago mensual" in response.data
    assert b"USD" in response.data
    assert b"75.50" in response.data


def test_recibo_muestra_numero_deterministico(app, db, base_data):
    pago = crear_pago(base_data)

    client = app.test_client()
    login(client, base_data["admin_a"])

    response = client.get(
        f"/finanzas/pagos/{pago.id}/recibo"
    )

    assert response.status_code == 200

    numero_esperado = f"REC-{pago.id:06d}".encode()
    assert numero_esperado in response.data


def test_recibo_pago_sin_aplicaciones(app, db, base_data):
    pago = crear_pago(
        base_data,
        valor="50.00",
    )

    client = app.test_client()
    login(client, base_data["admin_a"])

    response = client.get(
        f"/finanzas/pagos/{pago.id}/recibo"
    )

    assert response.status_code == 200

    assert (
        b"Este pago todav\xc3\xada no tiene aplicaciones."
        in response.data
    )

    assert b"50.00" in response.data
    assert b"0.00" in response.data


def test_recibo_muestra_aplicaciones(app, db, base_data):
    obligacion = crear_obligacion(
        base_data,
        valor="60.00",
        periodo="2026-09",
        concepto="Pension septiembre",
    )

    pago = crear_pago(
        base_data,
        valor="60.00",
        medio_pago="TRANSFERENCIA",
        referencia="BANCO-2026-001",
    )

    aplicar_pago_a_obligaciones(
        academia_id=pago.academia_id,
        pago_id=pago.id,
        aplicaciones=[
            {
                "obligacion_financiera_id": obligacion.id,
                "valor_aplicado": Decimal("40.00"),
            }
        ],
    )

    _db.session.commit()

    client = app.test_client()
    login(client, base_data["admin_a"])

    response = client.get(
        f"/finanzas/pagos/{pago.id}/recibo"
    )

    assert response.status_code == 200
    assert b"2026-09" in response.data
    assert b"PENSION" in response.data
    assert b"Pension septiembre" in response.data
    assert b"40.00" in response.data
    assert b"20.00" in response.data


def test_pago_anulado_conserva_recibo(app, db, base_data):
    pago = crear_pago(
        base_data,
        valor="35.00",
    )

    anular_pago(
        academia_id=pago.academia_id,
        pago_id=pago.id,
    )
    _db.session.commit()

    client = app.test_client()
    login(client, base_data["admin_a"])

    response = client.get(
        f"/finanzas/pagos/{pago.id}/recibo"
    )

    assert response.status_code == 200
    assert b"PAGO ANULADO" in response.data
    assert b"ANULADO" in response.data

    numero = f"REC-{pago.id:06d}".encode()
    assert numero in response.data


def test_otro_tenant_no_puede_ver_recibo(app, db, base_data):
    pago = crear_pago(
        base_data,
        academia_key="academia_b",
        alumno_key="alumno_b1",
    )

    client = app.test_client()
    login(client, base_data["admin_a"])

    response = client.get(
        f"/finanzas/pagos/{pago.id}/recibo"
    )

    assert response.status_code == 404


def test_profesor_puede_ver_recibo_alumno_visible(
    app,
    db,
    base_data,
):
    pago = crear_pago(base_data)

    client = app.test_client()
    login(client, base_data["profesor_a"])

    response = client.get(
        f"/finanzas/pagos/{pago.id}/recibo"
    )

    assert response.status_code == 200

    numero = f"REC-{pago.id:06d}".encode()
    assert numero in response.data


def test_detalle_pago_enlaza_recibo(app, db, base_data):
    pago = crear_pago(base_data)

    client = app.test_client()
    login(client, base_data["admin_a"])

    response = client.get(
        f"/finanzas/pagos/{pago.id}"
    )

    assert response.status_code == 200
    assert b"Ver recibo" in response.data

    ruta = (
        f"/finanzas/pagos/{pago.id}/recibo"
    ).encode()

    assert ruta in response.data


def test_recibo_indica_documento_no_tributario(
    app,
    db,
    base_data,
):
    pago = crear_pago(base_data)

    client = app.test_client()
    login(client, base_data["admin_a"])

    response = client.get(
        f"/finanzas/pagos/{pago.id}/recibo"
    )

    assert response.status_code == 200

    assert (
        b"No constituye factura ni comprobante tributario."
        in response.data
    )
def test_reporte_pagos_vacio(db, base_data):
    resumen, filas = obtener_reporte_pagos(
        academia_id=base_data["academia_a"].id,
    )

    assert filas == []
    assert resumen.cantidad_pagos == 0
    assert resumen.cantidad_registrados == 0
    assert resumen.cantidad_anulados == 0
    assert resumen.total_recibido == Decimal("0.00")
    assert resumen.total_aplicado == Decimal("0.00")
    assert resumen.saldo_sin_aplicar == Decimal("0.00")


def test_reporte_pagos_calcula_totales(db, base_data):
    pago_1 = crear_pago(
        base_data,
        valor="60.00",
        medio_pago="TRANSFERENCIA",
        referencia="TRX-001",
    )

    crear_pago(
        base_data,
        valor="40.00",
        medio_pago="EFECTIVO",
        referencia="CAJA-001",
    )

    obligacion = crear_obligacion(
        base_data,
        valor="60.00",
        concepto="Pension septiembre",
    )

    aplicar_pago_a_obligaciones(
        academia_id=pago_1.academia_id,
        pago_id=pago_1.id,
        aplicaciones=[
            {
                "obligacion_financiera_id": obligacion.id,
                "valor_aplicado": Decimal("50.00"),
            }
        ],
    )

    _db.session.commit()

    resumen, filas = obtener_reporte_pagos(
        academia_id=base_data["academia_a"].id,
    )

    assert len(filas) == 2
    assert resumen.cantidad_pagos == 2
    assert resumen.cantidad_registrados == 2
    assert resumen.cantidad_anulados == 0
    assert resumen.total_recibido == Decimal("100.00")
    assert resumen.total_aplicado == Decimal("50.00")
    assert resumen.saldo_sin_aplicar == Decimal("50.00")


def test_reporte_pagos_anulados_no_suman_totales(db, base_data):
    crear_pago(
        base_data,
        valor="60.00",
    )

    pago_anulado = crear_pago(
        base_data,
        valor="40.00",
    )

    anular_pago(
        academia_id=pago_anulado.academia_id,
        pago_id=pago_anulado.id,
    )

    _db.session.commit()

    resumen, filas = obtener_reporte_pagos(
        academia_id=base_data["academia_a"].id,
    )

    assert len(filas) == 2
    assert resumen.cantidad_pagos == 2
    assert resumen.cantidad_registrados == 1
    assert resumen.cantidad_anulados == 1
    assert resumen.total_recibido == Decimal("60.00")


def test_reporte_pagos_filtra_medio_pago(db, base_data):
    crear_pago(
        base_data,
        valor="30.00",
        medio_pago="EFECTIVO",
    )

    crear_pago(
        base_data,
        valor="70.00",
        medio_pago="TRANSFERENCIA",
    )

    resumen, filas = obtener_reporte_pagos(
        academia_id=base_data["academia_a"].id,
        medio_pago="TRANSFERENCIA",
    )

    assert len(filas) == 1
    assert filas[0].medio_pago == "TRANSFERENCIA"
    assert resumen.total_recibido == Decimal("70.00")


def test_reporte_pagos_filtra_estado(db, base_data):
    crear_pago(
        base_data,
        valor="50.00",
    )

    pago_anulado = crear_pago(
        base_data,
        valor="25.00",
    )

    anular_pago(
        academia_id=pago_anulado.academia_id,
        pago_id=pago_anulado.id,
    )

    _db.session.commit()

    resumen, filas = obtener_reporte_pagos(
        academia_id=base_data["academia_a"].id,
        estado="ANULADO",
    )

    assert len(filas) == 1
    assert filas[0].estado == "ANULADO"
    assert resumen.cantidad_anulados == 1
    assert resumen.total_recibido == Decimal("0.00")


def test_reporte_pagos_filtra_alumno(db, base_data):
    crear_pago(
        base_data,
        alumno_key="alumno_a1",
        valor="40.00",
    )

    crear_pago(
        base_data,
        alumno_key="alumno_a2",
        valor="30.00",
    )

    termino = base_data["alumno_a1"].nombres

    resumen, filas = obtener_reporte_pagos(
        academia_id=base_data["academia_a"].id,
        q=termino,
    )

    assert len(filas) == 1
    assert filas[0].alumno_id == base_data["alumno_a1"].id
    assert resumen.total_recibido == Decimal("40.00")


def test_reporte_pagos_respeta_tenant(db, base_data):
    crear_pago(
        base_data,
        academia_key="academia_a",
        alumno_key="alumno_a1",
        valor="40.00",
    )

    crear_pago(
        base_data,
        academia_key="academia_b",
        alumno_key="alumno_b1",
        valor="90.00",
    )

    resumen, filas = obtener_reporte_pagos(
        academia_id=base_data["academia_a"].id,
    )

    assert len(filas) == 1
    assert filas[0].alumno_id == base_data["alumno_a1"].id
    assert resumen.total_recibido == Decimal("40.00")


def test_admin_puede_ver_reporte_pagos(app, db, base_data):
    pago = crear_pago(
        base_data,
        valor="75.00",
        medio_pago="TRANSFERENCIA",
        referencia="REPORTE-001",
    )

    client = app.test_client()
    login(client, base_data["admin_a"])

    response = client.get(
        "/finanzas/reportes/pagos"
    )

    assert response.status_code == 200
    assert b"Reporte financiero de pagos" in response.data
    assert b"REPORTE-001" in response.data
    assert b"75.00" in response.data

    ruta_recibo = (
        f"/finanzas/pagos/{pago.id}/recibo"
    ).encode()

    assert ruta_recibo in response.data


def test_reporte_pagos_filtros_web(app, db, base_data):
    crear_pago(
        base_data,
        valor="30.00",
        medio_pago="EFECTIVO",
        referencia="EF-001",
    )

    crear_pago(
        base_data,
        valor="70.00",
        medio_pago="TRANSFERENCIA",
        referencia="TR-001",
    )

    client = app.test_client()
    login(client, base_data["admin_a"])

    response = client.get(
        "/finanzas/reportes/pagos"
        "?medio_pago=TRANSFERENCIA"
    )

    assert response.status_code == 200
    assert b"TR-001" in response.data
    assert b"EF-001" not in response.data


def test_cartera_enlaza_reporte_pagos(app, db, base_data):
    client = app.test_client()
    login(client, base_data["admin_a"])

    response = client.get(
        "/finanzas/cartera"
    )

    assert response.status_code == 200
    assert b"Reporte de pagos" in response.data
    assert b"/finanzas/reportes/pagos" in response.data
def test_profesor_reporte_pagos_limita_a_sucursal(
    app,
    db,
    base_data,
):
    sucursal_otro = Sucursal(
        nombre="Sucursal pagos externa",
        academia_id=base_data["academia_a"].id,
        activo=True,
    )

    _db.session.add(sucursal_otro)
    _db.session.flush()

    base_data["alumno_a2"].sucursal_id = sucursal_otro.id
    _db.session.commit()

    crear_pago(
        base_data,
        alumno_key="alumno_a1",
        valor="60.00",
        referencia="VISIBLE-PROFESOR",
    )

    crear_pago(
        base_data,
        alumno_key="alumno_a2",
        valor="40.00",
        referencia="OCULTO-PROFESOR",
    )

    client = app.test_client()
    login(client, base_data["profesor_a"])

    response = client.get(
        "/finanzas/reportes/pagos"
    )

    assert response.status_code == 200
    assert b"VISIBLE-PROFESOR" in response.data
    assert b"OCULTO-PROFESOR" not in response.data


def test_admin_reporte_pagos_ve_toda_academia(
    app,
    db,
    base_data,
):
    sucursal_otro = Sucursal(
        nombre="Sucursal pagos admin",
        academia_id=base_data["academia_a"].id,
        activo=True,
    )

    _db.session.add(sucursal_otro)
    _db.session.flush()

    base_data["alumno_a2"].sucursal_id = sucursal_otro.id
    _db.session.commit()

    crear_pago(
        base_data,
        alumno_key="alumno_a1",
        valor="60.00",
        referencia="ADMIN-MATRIZ",
    )

    crear_pago(
        base_data,
        alumno_key="alumno_a2",
        valor="40.00",
        referencia="ADMIN-OTRA-SUCURSAL",
    )

    client = app.test_client()
    login(client, base_data["admin_a"])

    response = client.get(
        "/finanzas/reportes/pagos"
    )

    assert response.status_code == 200
    assert b"ADMIN-MATRIZ" in response.data
    assert b"ADMIN-OTRA-SUCURSAL" in response.data


def test_reporte_pagos_filtra_rango_fechas(
    app,
    db,
    base_data,
):
    crear_pago(
        base_data,
        valor="30.00",
        referencia="PAGO-SEPTIEMBRE",
        fecha_pago=date(2026, 9, 5),
    )

    crear_pago(
        base_data,
        valor="70.00",
        referencia="PAGO-OCTUBRE",
        fecha_pago=date(2026, 10, 10),
    )

    client = app.test_client()
    login(client, base_data["admin_a"])

    response = client.get(
        "/finanzas/reportes/pagos"
        "?fecha_desde=2026-10-01"
        "&fecha_hasta=2026-10-31"
    )

    assert response.status_code == 200
    assert b"PAGO-OCTUBRE" in response.data
    assert b"PAGO-SEPTIEMBRE" not in response.data


def test_reporte_pagos_rango_fechas_invalido(
    app,
    db,
    base_data,
):
    crear_pago(
        base_data,
        valor="50.00",
        referencia="NO-DEBE-MOSTRARSE",
    )

    client = app.test_client()
    login(client, base_data["admin_a"])

    response = client.get(
        "/finanzas/reportes/pagos"
        "?fecha_desde=2026-10-31"
        "&fecha_hasta=2026-10-01",
        follow_redirects=True,
    )

    assert response.status_code == 200
    assert (
        b"La fecha desde no puede ser posterior"
        in response.data
    )
    assert b"NO-DEBE-MOSTRARSE" not in response.data
def test_admin_puede_ver_reporte_cartera(
    app,
    db,
    base_data,
):
    crear_obligacion(
        base_data,
        alumno_key="alumno_a1",
        valor="60.00",
        concepto="Pension septiembre",
    )

    client = app.test_client()
    login(client, base_data["admin_a"])

    response = client.get(
        "/finanzas/reportes/cartera"
    )

    assert response.status_code == 200
    assert b"Reporte financiero de cartera" in response.data
    assert b"Antig\xc3\xbcedad de cartera" in response.data

    nombre = (
        f"{base_data['alumno_a1'].apellidos} "
        f"{base_data['alumno_a1'].nombres}"
    ).encode()

    assert nombre in response.data
    assert b"60.00" in response.data


def test_reporte_cartera_muestra_aging(
    app,
    db,
    base_data,
):
    crear_obligacion(
        base_data,
        valor="60.00",
    )

    client = app.test_client()
    login(client, base_data["admin_a"])

    response = client.get(
        "/finanzas/reportes/cartera"
    )

    assert response.status_code == 200
    assert b"1 - 30 d\xc3\xadas" in response.data
    assert b"31 - 60 d\xc3\xadas" in response.data
    assert b"61 - 90 d\xc3\xadas" in response.data
    assert b"M\xc3\xa1s de 90 d\xc3\xadas" in response.data
    assert b"Sin vencimiento" in response.data


def test_reporte_cartera_filtra_alumno_y_resumen(
    app,
    db,
    base_data,
):
    crear_obligacion(
        base_data,
        alumno_key="alumno_a1",
        valor="60.00",
    )

    crear_obligacion(
        base_data,
        alumno_key="alumno_a2",
        valor="40.00",
        concepto="Pension alumno dos",
    )

    client = app.test_client()
    login(client, base_data["admin_a"])

    termino = base_data["alumno_a1"].nombres

    response = client.get(
        "/finanzas/reportes/cartera",
        query_string={
            "q": termino,
        },
    )

    assert response.status_code == 200

    nombre_visible = (
        f"{base_data['alumno_a1'].apellidos} "
        f"{base_data['alumno_a1'].nombres}"
    ).encode()

    nombre_oculto = (
        f"{base_data['alumno_a2'].apellidos} "
        f"{base_data['alumno_a2'].nombres}"
    ).encode()

    assert nombre_visible in response.data
    assert nombre_oculto not in response.data

    # El resumen debe corresponder también
    # al resultado filtrado, no a toda la academia.
    assert b"60.00" in response.data
    assert b"100.00" not in response.data


def test_profesor_reporte_cartera_limita_a_sucursal(
    app,
    db,
    base_data,
):
    sucursal_otro = Sucursal(
        nombre="Sucursal cartera externa",
        academia_id=base_data["academia_a"].id,
        activo=True,
    )

    _db.session.add(sucursal_otro)
    _db.session.flush()

    base_data["alumno_a2"].sucursal_id = sucursal_otro.id
    _db.session.commit()

    crear_obligacion(
        base_data,
        alumno_key="alumno_a1",
        valor="60.00",
    )

    crear_obligacion(
        base_data,
        alumno_key="alumno_a2",
        valor="40.00",
        concepto="Pension sucursal externa",
    )

    client = app.test_client()
    login(client, base_data["profesor_a"])

    response = client.get(
        "/finanzas/reportes/cartera"
    )

    assert response.status_code == 200

    nombre_visible = (
        f"{base_data['alumno_a1'].apellidos} "
        f"{base_data['alumno_a1'].nombres}"
    ).encode()

    nombre_oculto = (
        f"{base_data['alumno_a2'].apellidos} "
        f"{base_data['alumno_a2'].nombres}"
    ).encode()

    assert nombre_visible in response.data
    assert nombre_oculto not in response.data
    assert b"40.00" not in response.data


def test_cartera_enlaza_reporte_cartera(
    app,
    db,
    base_data,
):
    client = app.test_client()
    login(client, base_data["admin_a"])

    response = client.get(
        "/finanzas/cartera"
    )

    assert response.status_code == 200
    assert b"Reporte de cartera" in response.data
    assert b"/finanzas/reportes/cartera" in response.data

def test_exportar_excel_pagos_genera_xlsx_valido(
    app,
    db,
    base_data,
):
    crear_pago(
        base_data,
        valor="75.50",
        medio_pago="TRANSFERENCIA",
        referencia="EXCEL-PAGO-001",
    )

    client = app.test_client()
    login(client, base_data["admin_a"])

    response = client.get(
        "/finanzas/reportes/pagos/excel"
    )

    assert response.status_code == 200
    assert (
        response.mimetype
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert ".xlsx" in response.headers["Content-Disposition"]

    libro = load_workbook(
        BytesIO(response.data),
        data_only=True,
    )

    assert libro.sheetnames == ["Pagos"]

    hoja = libro["Pagos"]

    encabezados = [
        celda.value
        for celda in hoja[1]
    ]

    assert encabezados == [
        "Fecha",
        "Alumno",
        "Identificación",
        "Medio de pago",
        "Referencia",
        "Estado",
        "Moneda",
        "Valor",
        "Aplicado",
        "Saldo disponible",
    ]

    assert hoja.max_row == 2

    assert hoja["D2"].value == "TRANSFERENCIA"
    assert hoja["E2"].value == "EXCEL-PAGO-001"
    assert hoja["F2"].value == "REGISTRADO"
    assert hoja["G2"].value == "USD"
    assert hoja["H2"].value == 75.5


def test_exportar_excel_pagos_respeta_filtros(
    app,
    db,
    base_data,
):
    crear_pago(
        base_data,
        valor="30.00",
        medio_pago="EFECTIVO",
        referencia="EXCEL-EFECTIVO",
    )

    crear_pago(
        base_data,
        valor="70.00",
        medio_pago="TRANSFERENCIA",
        referencia="EXCEL-TRANSFERENCIA",
    )

    client = app.test_client()
    login(client, base_data["admin_a"])

    response = client.get(
        "/finanzas/reportes/pagos/excel",
        query_string={
            "medio_pago": "TRANSFERENCIA",
        },
    )

    assert response.status_code == 200

    libro = load_workbook(
        BytesIO(response.data),
        data_only=True,
    )

    hoja = libro["Pagos"]

    referencias = [
        hoja.cell(
            row=fila,
            column=5,
        ).value
        for fila in range(
            2,
            hoja.max_row + 1,
        )
    ]

    assert "EXCEL-TRANSFERENCIA" in referencias
    assert "EXCEL-EFECTIVO" not in referencias


def test_profesor_excel_pagos_limita_a_sucursal(
    app,
    db,
    base_data,
):
    sucursal_otro = Sucursal(
        nombre="Sucursal excel pagos externa",
        academia_id=base_data["academia_a"].id,
        activo=True,
    )

    _db.session.add(sucursal_otro)
    _db.session.flush()

    base_data["alumno_a2"].sucursal_id = (
        sucursal_otro.id
    )

    _db.session.commit()

    crear_pago(
        base_data,
        alumno_key="alumno_a1",
        valor="60.00",
        referencia="EXCEL-PAGO-VISIBLE",
    )

    crear_pago(
        base_data,
        alumno_key="alumno_a2",
        valor="40.00",
        referencia="EXCEL-PAGO-OCULTO",
    )

    client = app.test_client()
    login(client, base_data["profesor_a"])

    response = client.get(
        "/finanzas/reportes/pagos/excel"
    )

    assert response.status_code == 200

    libro = load_workbook(
        BytesIO(response.data),
        data_only=True,
    )

    hoja = libro["Pagos"]

    referencias = [
        hoja.cell(
            row=fila,
            column=5,
        ).value
        for fila in range(
            2,
            hoja.max_row + 1,
        )
    ]

    assert "EXCEL-PAGO-VISIBLE" in referencias
    assert "EXCEL-PAGO-OCULTO" not in referencias


def test_exportar_excel_cartera_genera_resumen_y_detalle(
    app,
    db,
    base_data,
):
    crear_obligacion(
        base_data,
        alumno_key="alumno_a1",
        valor="60.00",
    )

    client = app.test_client()
    login(client, base_data["admin_a"])

    response = client.get(
        "/finanzas/reportes/cartera/excel"
    )

    assert response.status_code == 200
    assert (
        response.mimetype
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert ".xlsx" in response.headers["Content-Disposition"]

    libro = load_workbook(
        BytesIO(response.data),
        data_only=True,
    )

    assert libro.sheetnames == [
        "Resumen",
        "Cartera",
    ]

    resumen = libro["Resumen"]
    cartera = libro["Cartera"]

    indicadores = {
        resumen.cell(
            row=fila,
            column=1,
        ).value:
        resumen.cell(
            row=fila,
            column=2,
        ).value
        for fila in range(
            2,
            resumen.max_row + 1,
        )
    }

    assert indicadores["Total obligaciones"] == 60
    assert indicadores["Saldo pendiente"] == 60
    assert indicadores["Saldo vencido"] == 60

    encabezados = [
        celda.value
        for celda in cartera[1]
    ]

    assert encabezados == [
        "Alumno",
        "Total obligaciones",
        "Total aplicado",
        "Saldo pendiente",
        "Saldo vigente",
        "Saldo vencido",
        "Días atraso",
        "Saldo disponible",
        "Pendientes",
        "Parciales",
        "Vencidas",
        "Último pago",
    ]

    assert cartera.max_row == 3

    nombres = [
        cartera.cell(
            row=fila,
            column=1,
        ).value
        for fila in range(
            2,
            cartera.max_row + 1,
        )
    ]

    nombre_esperado = (
        f"{base_data['alumno_a1'].apellidos} "
        f"{base_data['alumno_a1'].nombres}"
    )

    assert nombre_esperado in nombres


def test_exportar_excel_cartera_respeta_filtro_alumno(
    app,
    db,
    base_data,
):
    crear_obligacion(
        base_data,
        alumno_key="alumno_a1",
        valor="60.00",
    )

    crear_obligacion(
        base_data,
        alumno_key="alumno_a2",
        valor="40.00",
        concepto="Pension excel alumno dos",
    )

    client = app.test_client()
    login(client, base_data["admin_a"])

    response = client.get(
        "/finanzas/reportes/cartera/excel",
        query_string={
            "q": base_data["alumno_a1"].nombres,
        },
    )

    assert response.status_code == 200

    libro = load_workbook(
        BytesIO(response.data),
        data_only=True,
    )

    resumen = libro["Resumen"]
    cartera = libro["Cartera"]

    nombres = [
        cartera.cell(
            row=fila,
            column=1,
        ).value
        for fila in range(
            2,
            cartera.max_row + 1,
        )
    ]

    nombre_visible = (
        f"{base_data['alumno_a1'].apellidos} "
        f"{base_data['alumno_a1'].nombres}"
    )

    nombre_oculto = (
        f"{base_data['alumno_a2'].apellidos} "
        f"{base_data['alumno_a2'].nombres}"
    )

    assert nombre_visible in nombres
    assert nombre_oculto not in nombres

    indicadores = {
        resumen.cell(
            row=fila,
            column=1,
        ).value:
        resumen.cell(
            row=fila,
            column=2,
        ).value
        for fila in range(
            2,
            resumen.max_row + 1,
        )
    }

    assert indicadores["Total obligaciones"] == 60
    assert indicadores["Saldo pendiente"] == 60


def test_profesor_excel_cartera_limita_a_sucursal(
    app,
    db,
    base_data,
):
    sucursal_otro = Sucursal(
        nombre="Sucursal excel cartera externa",
        academia_id=base_data["academia_a"].id,
        activo=True,
    )

    _db.session.add(sucursal_otro)
    _db.session.flush()

    base_data["alumno_a2"].sucursal_id = (
        sucursal_otro.id
    )

    _db.session.commit()

    crear_obligacion(
        base_data,
        alumno_key="alumno_a1",
        valor="60.00",
    )

    crear_obligacion(
        base_data,
        alumno_key="alumno_a2",
        valor="40.00",
        concepto="Pension excel sucursal externa",
    )

    client = app.test_client()
    login(client, base_data["profesor_a"])

    response = client.get(
        "/finanzas/reportes/cartera/excel"
    )

    assert response.status_code == 200

    libro = load_workbook(
        BytesIO(response.data),
        data_only=True,
    )

    cartera = libro["Cartera"]

    nombres = [
        cartera.cell(
            row=fila,
            column=1,
        ).value
        for fila in range(
            2,
            cartera.max_row + 1,
        )
    ]

    nombre_visible = (
        f"{base_data['alumno_a1'].apellidos} "
        f"{base_data['alumno_a1'].nombres}"
    )

    nombre_oculto = (
        f"{base_data['alumno_a2'].apellidos} "
        f"{base_data['alumno_a2'].nombres}"
    )

    assert nombre_visible in nombres
    assert nombre_oculto not in nombres
def test_profesor_no_puede_ver_recibo_de_otra_sucursal(
    app,
    db,
    base_data,
):
    sucursal_otro = Sucursal(
        nombre="Sucursal recibo externa",
        academia_id=base_data["academia_a"].id,
        activo=True,
    )

    _db.session.add(sucursal_otro)
    _db.session.flush()

    base_data["alumno_a2"].sucursal_id = sucursal_otro.id
    _db.session.commit()

    pago = crear_pago(
        base_data,
        alumno_key="alumno_a2",
        valor="40.00",
        referencia="RECIBO-OTRA-SUCURSAL",
    )

    client = app.test_client()
    login(client, base_data["profesor_a"])

    response = client.get(
        f"/finanzas/pagos/{pago.id}/recibo"
    )

    assert response.status_code == 403


def test_rol_sin_finanzas_no_accede_a_reportes_b7(
    app,
    db,
    base_data,
):
    role = Role.query.filter_by(
        name="COACH"
    ).first()

    if role is None:
        role = Role(
            name="COACH",
            description="COACH",
        )
        _db.session.add(role)
        _db.session.flush()

    usuario = base_data["admin_a"]
    usuario.roles = [role]
    _db.session.commit()

    client = app.test_client()
    login(client, usuario)

    endpoints = [
        "/finanzas/reportes/pagos",
        "/finanzas/reportes/pagos/excel",
        "/finanzas/reportes/cartera",
        "/finanzas/reportes/cartera/excel",
    ]

    for endpoint in endpoints:
        response = client.get(endpoint)
        assert response.status_code == 403


def test_superadmin_con_academia_accede_a_reportes_b7(
    app,
    db,
    base_data,
):
    role = Role.query.filter_by(
        name="SUPERADMIN"
    ).first()

    if role is None:
        role = Role(
            name="SUPERADMIN",
            description="Acceso total",
        )
        _db.session.add(role)
        _db.session.flush()

    usuario = base_data["admin_a"]
    usuario.roles = [role]
    _db.session.commit()

    client = app.test_client()
    login(client, usuario)

    endpoints = [
        "/finanzas/reportes/pagos",
        "/finanzas/reportes/pagos/excel",
        "/finanzas/reportes/cartera",
        "/finanzas/reportes/cartera/excel",
    ]

    for endpoint in endpoints:
        response = client.get(endpoint)
        assert response.status_code == 200


def test_reportes_b7_respetan_tenant_en_excel(
    app,
    db,
    base_data,
):
    crear_pago(
        base_data,
        academia_key="academia_a",
        alumno_key="alumno_a1",
        valor="60.00",
        referencia="TENANT-A-PAGO",
    )

    crear_pago(
        base_data,
        academia_key="academia_b",
        alumno_key="alumno_b1",
        valor="900.00",
        referencia="TENANT-B-PAGO",
    )

    crear_obligacion(
        base_data,
        academia_key="academia_a",
        alumno_key="alumno_a1",
        valor="60.00",
        concepto="TENANT-A-CARTERA",
    )

    crear_obligacion(
        base_data,
        academia_key="academia_b",
        alumno_key="alumno_b1",
        valor="900.00",
        concepto="TENANT-B-CARTERA",
    )

    client = app.test_client()
    login(client, base_data["admin_a"])

    pagos_response = client.get(
        "/finanzas/reportes/pagos/excel"
    )

    assert pagos_response.status_code == 200

    libro_pagos = load_workbook(
        BytesIO(pagos_response.data),
        data_only=True,
    )

    hoja_pagos = libro_pagos["Pagos"]

    referencias = [
        hoja_pagos.cell(
            row=fila,
            column=5,
        ).value
        for fila in range(
            2,
            hoja_pagos.max_row + 1,
        )
    ]

    assert "TENANT-A-PAGO" in referencias
    assert "TENANT-B-PAGO" not in referencias

    cartera_response = client.get(
        "/finanzas/reportes/cartera/excel"
    )

    assert cartera_response.status_code == 200

    libro_cartera = load_workbook(
        BytesIO(cartera_response.data),
        data_only=True,
    )

    hoja_cartera = libro_cartera["Cartera"]

    nombres = [
        hoja_cartera.cell(
            row=fila,
            column=1,
        ).value
        for fila in range(
            2,
            hoja_cartera.max_row + 1,
        )
    ]

    nombre_a = (
        f"{base_data['alumno_a1'].apellidos} "
        f"{base_data['alumno_a1'].nombres}"
    )

    nombre_b = (
        f"{base_data['alumno_b1'].apellidos} "
        f"{base_data['alumno_b1'].nombres}"
    )

    assert nombre_a in nombres
    assert nombre_b not in nombres


def test_reportes_b7_enlazan_exportacion_excel(
    app,
    db,
    base_data,
):
    client = app.test_client()
    login(client, base_data["admin_a"])

    pagos = client.get(
        "/finanzas/reportes/pagos"
    )

    cartera = client.get(
        "/finanzas/reportes/cartera"
    )

    assert pagos.status_code == 200
    assert cartera.status_code == 200

    assert (
        b"/finanzas/reportes/pagos/excel"
        in pagos.data
    )

    assert (
        b"/finanzas/reportes/cartera/excel"
        in cartera.data
    )