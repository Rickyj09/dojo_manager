from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from app.models.grado import Grado
from app.models.participacion import Participacion
from app.models.sucursal import Sucursal
from app.models.torneo import Torneo
from test_finanzas_tarifarios_ui import cliente


def _texto_excel(response):
    workbook = load_workbook(
        BytesIO(response.data),
        read_only=True,
        data_only=True,
    )

    textos = []

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            for value in row:
                if value is not None:
                    textos.append(str(value))

    return "\n".join(textos)


def test_admin_reporte_general_solo_ve_su_academia(
    app,
    db,
    base_data,
):
    base_data["alumno_a1"].nombres = "ALUMNO_A_VISIBLE"
    base_data["alumno_b1"].nombres = "ALUMNO_B_PROHIBIDO"
    db.session.commit()

    client = cliente(
        app,
        base_data["admin_a"],
    )

    response = client.get("/reportes/")

    assert response.status_code == 200
    assert "ALUMNO_A_VISIBLE" in response.text
    assert "ALUMNO_B_PROHIBIDO" not in response.text


def test_admin_no_puede_forzar_sucursal_otra_academia(
    app,
    db,
    base_data,
):
    base_data["alumno_b1"].nombres = "ALUMNO_B_PROHIBIDO"
    db.session.commit()

    client = cliente(
        app,
        base_data["admin_a"],
    )

    response = client.get(
        "/reportes/",
        query_string={
            "sucursal_id":
                base_data["alumno_b1"].sucursal_id,
        },
    )

    assert response.status_code == 200
    assert "ALUMNO_B_PROHIBIDO" not in response.text


def test_excel_general_solo_contiene_su_academia(
    app,
    db,
    base_data,
):
    base_data["alumno_a1"].nombres = "EXCEL_ALUMNO_A"
    base_data["alumno_b1"].nombres = "EXCEL_ALUMNO_B"
    db.session.commit()

    client = cliente(
        app,
        base_data["admin_a"],
    )

    response = client.get(
        "/reportes/export/excel"
    )

    assert response.status_code == 200

    contenido = _texto_excel(response)

    assert "EXCEL_ALUMNO_A" in contenido
    assert "EXCEL_ALUMNO_B" not in contenido

def test_morosidad_solo_muestra_su_academia(
    app,
    db,
    base_data,
):
    base_data["alumno_a1"].nombres = "MOROSO_A_VISIBLE"
    base_data["alumno_b1"].nombres = "MOROSO_B_PROHIBIDO"
    db.session.commit()

    client = cliente(
        app,
        base_data["admin_a"],
    )

    response = client.get(
        "/reportes/morosidad",
        query_string={
            "solo_morosos": "0",
            "activo": "1",
        },
    )

    assert response.status_code == 200
    assert "MOROSO_A_VISIBLE" in response.text
    assert "MOROSO_B_PROHIBIDO" not in response.text


def test_grados_del_combo_solo_son_de_su_academia(
    app,
    db,
    base_data,
):
    grado_a = Grado(
        academia_id=base_data["academia_a"].id,
        nombre="GRADO_TENANT_A",
        tipo="KUP",
        orden=901,
        activo=True,
    )

    grado_b = Grado(
        academia_id=base_data["academia_b"].id,
        nombre="GRADO_TENANT_B",
        tipo="KUP",
        orden=902,
        activo=True,
    )

    db.session.add_all([
        grado_a,
        grado_b,
    ])
    db.session.commit()

    client = cliente(
        app,
        base_data["admin_a"],
    )

    response = client.get("/reportes/")

    assert response.status_code == 200
    assert "GRADO_TENANT_A" in response.text
    assert "GRADO_TENANT_B" not in response.text


def test_lista_torneos_solo_muestra_su_academia(
    app,
    db,
    base_data,
):
    torneo_a = Torneo(
        academia_id=base_data["academia_a"].id,
        nombre="TORNEO_TENANT_A",
        fecha=date(2026, 10, 10),
        activo=True,
    )

    torneo_b = Torneo(
        academia_id=base_data["academia_b"].id,
        nombre="TORNEO_TENANT_B",
        fecha=date(2026, 10, 11),
        activo=True,
    )

    db.session.add_all([
        torneo_a,
        torneo_b,
    ])
    db.session.commit()

    client = cliente(
        app,
        base_data["admin_a"],
    )

    response = client.get(
        "/reportes/seleccion"
    )

    assert response.status_code == 200
    assert "TORNEO_TENANT_A" in response.text
    assert "TORNEO_TENANT_B" not in response.text


def test_admin_no_puede_abrir_torneo_otra_academia(
    app,
    db,
    base_data,
):
    torneo_b = Torneo(
        academia_id=base_data["academia_b"].id,
        nombre="TORNEO_PRIVADO_B",
        fecha=date(2026, 10, 12),
        activo=True,
    )

    db.session.add(torneo_b)
    db.session.commit()

    client = cliente(
        app,
        base_data["admin_a"],
    )

    response = client.get(
        f"/reportes/torneo/{torneo_b.id}/seleccionar"
    )

    assert response.status_code == 404


def test_admin_no_puede_exportar_torneo_otra_academia(
    app,
    db,
    base_data,
):
    torneo_b = Torneo(
        academia_id=base_data["academia_b"].id,
        nombre="TORNEO_EXPORT_PRIVADO_B",
        fecha=date(2026, 10, 13),
        activo=True,
    )

    db.session.add(torneo_b)
    db.session.commit()

    client = cliente(
        app,
        base_data["admin_a"],
    )

    response = client.get(
        f"/reportes/torneo/{torneo_b.id}/seleccion.xlsx"
    )

    assert response.status_code == 404


def test_profesor_reporte_respeta_sucursal(
    app,
    db,
    base_data,
):
    otra_sucursal = Sucursal(
        academia_id=base_data["academia_a"].id,
        nombre="Sucursal Profesor No Visible",
        activo=True,
    )

    db.session.add(otra_sucursal)
    db.session.flush()

    base_data["alumno_a1"].nombres = "PROFESOR_VISIBLE"

    base_data["alumno_a2"].nombres = "PROFESOR_NO_VISIBLE"
    base_data["alumno_a2"].sucursal_id = otra_sucursal.id

    db.session.commit()

    client = cliente(
        app,
        base_data["profesor_a"],
    )

    response = client.get("/reportes/")

    assert response.status_code == 200
    assert "PROFESOR_VISIBLE" in response.text
    assert "PROFESOR_NO_VISIBLE" not in response.text
def test_post_manipulado_no_selecciona_alumno_de_otro_tenant(
    app,
    db,
    base_data,
):
    torneo = Torneo(
        academia_id=base_data["academia_a"].id,
        nombre="TORNEO_POST_TENANT_A",
        fecha=date(2026, 10, 20),
        activo=True,
    )

    db.session.add(torneo)
    db.session.commit()

    alumno_ajeno = base_data["alumno_b1"]

    client = cliente(
        app,
        base_data["admin_a"],
    )

    response = client.post(
        f"/reportes/torneo/{torneo.id}/seleccionar",
        data={
            "alumno_ids[]": str(alumno_ajeno.id),
            f"modalidad_{alumno_ajeno.id}": "COMBATE",
        },
        follow_redirects=False,
    )

    assert response.status_code == 302

    assert (
        Participacion.query
        .filter_by(
            torneo_id=torneo.id,
            alumno_id=alumno_ajeno.id,
        )
        .count()
        == 0
    )